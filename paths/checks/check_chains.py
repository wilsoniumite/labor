# check_chains.py — the machine chain in the data (code/chains.py): the gate, the accounting identities, the orderings the
# conversions must respect, the hours source against the lambda thread's, the full-chain cost split of calibration (c)'s
# bundle, one independent path from the raw BEA files (no lambda-thread engine, no chains.py code), the model against the
# data from macro.py's own calibrations, and every headline number read back from the results the module wrote.
#
# THRESHOLDS SET AFTER SEEING THE DATA (they describe what was found; they are not a priori bounds):
#   C2  the detail reconstruction tolerance 1e-3 (observed 6.2e-4; the summary keeps the thread's a-priori 5e-4);
#   C6  the cap's excess under 2% of the I-O's GOS every year (observed at most 1.53%);
#   C12 the FTE basis above the thread's hours-worked basis by 4-20% (observed 4.8-19.1%);
#   C15 detail against summary within 0.02 (observed 0.011);
#   C20 the self-built (Use/Make levels) domp recompute within 1e-5 (the published DR's 7-digit rounding; the Unit 1
#       verifier had observed 2.7e-6 on this route before the tolerance was set; observed here 2.5e-7).
# Every other tolerance is numerical (float rounding, the published TR's 7-digit rounding) and was set before looking.
#
# Needs BEA's files in the gitignored progress_and_prosperity/lambda/data/cache (a fresh clone: code/chains.py --pull
# fetches them); without them it stops at once and says so.
#
# Run from paths/:  ../venv/Scripts/python.exe checks/check_chains.py   (clean under -W error: every file is closed)
# Exit status 0 only if every check passes; prints "ALL GREEN (n)" or the failures.

from __future__ import annotations

import csv
import io
import json
import os
import re
import sys
import zipfile

import numpy as np
import pandas as pd

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
REPO = os.path.dirname(ROOT)
CACHE = os.path.join(REPO, "progress_and_prosperity", "lambda", "data", "cache")
NEEDS = ("AllTablesIO.zip", "AllTablesSUP.zip", "Section2All_xls.xlsx", "Section3All_xls.xlsx", "ValueAdded.xlsx", "NipaDataA.txt")
MISSING = [n for n in NEEDS if not os.path.isfile(os.path.join(CACHE, n))]
if MISSING:
    raise SystemExit(f"BEA's files are missing from the gitignored cache {CACHE}: {', '.join(MISSING)}. Fetch them first, from "
                     "paths/: ../venv/Scripts/python.exe code/chains.py --pull (the I-O zips also come with the lambda "
                     "thread's progress_and_prosperity/lambda/code/pull_family_a.py)")
sys.path.insert(0, os.path.join(ROOT, "code"))
import macro as m  # noqa: E402  (the model side of C21; chains.py is not imported anywhere here)

RESULTS = []


def check(name, ok, detail=""):
    RESULTS.append((name, bool(ok), detail))
    print(f"  [{'PASS' if ok else 'FAIL'}] {name}" + (f" — {detail}" if detail else ""))


with open(os.path.join(ROOT, "results", "chains.json"), encoding="utf-8") as fh:
    R = json.load(fh)
REF = pd.read_csv(os.path.join(REPO, "progress_and_prosperity", "lambda", "data", "lambda_us_family_a.csv")).set_index("year")
YEARS = [str(y) for y in range(1997, 2024)]
TYPES = list(R["types"])
AGG = R["aggregates"]
SET = "machinery_software_set"
BUNDLE_C = "equipment_ipp_structures"
BASES = ("tot", "dom", "domp")
A = AGG[f"{SET}|investment"]
PCTS = ["P10", "P25", "P50", "P75", "P90"]
M = lambda xs: float(np.mean(xs))  # noqa: E731
W15 = [str(y) for y in range(2015, 2024)]
RC = str(R["meta"]["rho_central"])
SVC, SVT = f"service_rho_{RC}", f"service_traced_rho_{RC}"

# ---- 1. the gate
g = R["gate"]
dev = max(abs(AGG[f"{s}|thread"][y][f"theta_{t}"] - REF.loc[int(y), f"lam_{s}_{t}"])
          for s in ("narrow", "medium", "broad") for t in ("tot", "dom", "domp") for y in YEARS)
check("C1 gate: the lambda thread's narrow/medium/broad tot/dom/domp series (lambda_us_family_a.csv) reproduced to <= 1e-6 "
      "in every year 1997-2023, both by the module's own gate and read back from its stored aggregates",
      g["max_abs_dev"] <= 1e-6 and g["n_values"] == 27 * 9 and dev <= 1e-6, f"module {g['max_abs_dev']:.1e}, read back {dev:.1e}")
D = R["detail_2017"]
check("C2 reconstruction: W(I - BW)^-1 reproduces the published IxC total requirements to <= 5e-4 at summary every year (the "
      "thread's tolerance) and to <= 1e-3 at the 2017 detail (402 industries; DR rounding) [the detail tolerance was set "
      "after seeing 6.2e-4], and at detail value added plus the inputs with no domestic output (noncomparable imports, used "
      "goods) resolves to 1 for every produced commodity",
      g["recon_err_max"] <= 5e-4 and D["recon_err"] <= D["recon_tol"] and D["va_resolution_max_dev"] <= 1e-4
      and D["industries"] == 402,
      f"summary {g['recon_err_max']:.1e}; detail {D['recon_err']:.1e}; detail VA identity {D['va_resolution_max_dev']:.1e}")

# ---- 2. identities
rows = [(f"type {c}", R["types"][c]["annual"][y]) for c in TYPES for y in YEARS] + \
       [(k, v[y]) for k, v in AGG.items() for y in YEARS]
e_tot = max(abs(s["theta_tot"] + s["tax_tot"] + s["gos_tot"] - 1.0) for _, s in rows)
check("C3 share identity (tot): compensation + taxes on production less subsidies + gross operating surplus = 1 for every "
      "type, aggregate and year", e_tot <= 1e-5, f"max dev {e_tot:.1e}")
e_dom = max(abs(s["theta_dom"] + s["tax_dom"] + s["gos_dom"] + s["imports_dom"] - 1.0) for _, s in rows)
check("C4 domestic leakage identity: compensation + taxes + GOS + imported intermediates = 1 on the domestic system (dom)",
      e_dom <= 1e-5, f"max dev {e_dom:.1e}")
e_cfc = max(abs(s[f"cfc_{b}{v}"] + s[f"nos_{b}{v}"] - s[f"gos_{b}"]) for _, s in rows
            for v in ("", "_uncapped", "_economywide") for b in BASES)
led = R["cfc_ledger"]
cov = max(abs(led[y]["fa_mapped_private"] / led[y]["fa_private_total"] - 1.0) for y in YEARS)
gg = max(abs(led[y]["general_gov_gos_gdpxind"] / led[y]["general_gov_cfc_nipa"] - 1.0) for y in YEARS)
check("C5 CFC + NOS = GOS in every variant and on every basis; the industry map covers BEA's private fixed assets exactly "
      "once (mapped CFC = FA total private) and general government's GOS equals its CFC (GDP by industry against NIPA 7.5), "
      "every year", e_cfc <= 1e-12 and cov <= 1e-4 and gg <= 1e-4,
      f"identity {e_cfc:.1e}; FA coverage {cov:.1e}; general government {gg:.1e}")
cap = max(led[y]["cap_excess_share_of_io_gos"] for y in YEARS)
binds334 = [y for y in YEARS if "334" in led[y]["cap_binds"]]
CF = R["cfc_cap_and_fallback"]
n_capped = sum(len(led[y]["cap_binds"]) for y in YEARS)
n_fb = sum(len(led[y]["per_output_fallback"]) for y in YEARS)
detail6 = (f"max excess {cap:.4f} of GOS; {n_capped} capped and {n_fb} fallback industry-years; chain-weighted, per $1 of "
           f"final output of the machinery + software set (the (c) bundle in cfc_cap_and_fallback):\n        "
           + "\n        ".join(R["notes"]["cfc_cap_and_fallback"][f"per_industry_{SET}"]))
check("C6 finding: the FA-by-industry and GDP-by-industry bases do not match cleanly where IPP (R&D) is large — FA's CFC "
      "exceeds GOS for 334 in some years (the cap at GOS binds) — but the excess the cap removes stays under 2% of the I-O's "
      "GOS every year [threshold set after seeing the data: observed 1.53%]; every capped and fallback industry-year is "
      "listed with its chain-weighted contribution (5415, 22-24% of the set's weight in its capped and fallback years, among "
      "them); the uncapped and "
      "economy-wide-ratio variants are reported beside", cap < 0.02 and len(binds334) > 0, detail6)

# ---- 3. orderings
viol = []
for c in TYPES:
    t = R["types"][c]
    for y in YEARS:
        th = t["annual"][y]["theta_tot"]
        for r, v in t["service"].items():
            if not v[y] <= th + 1e-12:
                viol.append((c, y, r))
        if not th <= t["traced"][y] + 1e-12:
            viol.append((c, y, "traced"))
S = R[f"{SET}_service"]
for y in YEARS:
    th = A[y]["theta_tot"]
    for r, v in S["service"].items():
        if not v[y] <= th + 1e-12:
            viol.append(("agg", y, r))
    if not th <= S["traced"][y] + 1e-12:
        viol.append(("agg", y, "traced"))
    if not S["service"]["0.08"][y] <= S["service"][RC][y] <= S["service"]["0.04"][y]:
        viol.append(("agg", y, "rho order"))
check("C7 theta_service <= theta_asset <= theta+ for every type, the aggregate, every year and every rho; the service share "
      "falls as rho rises", not viol, f"violations {viol[:5]}")
dd = max(max(s["theta_dom"] - s["theta_tot"], s["theta_domp"] - s["theta_dom"]) for _, s in rows)
check("C8 the import purge only removes labour: theta_domp <= theta_dom <= theta_tot everywhere (tolerance 1e-4: the "
      "published TR against the reconstruction)", dd <= 1e-4, f"max excess {dd:.1e}")
ok_w = all(abs(sum(A[y]["weights"].values()) - 1.0) < 1e-12 for y in YEARS) and S["consistency_max_dev"] <= 1e-9
check("C9 the investment weights sum to 1 each year and the aggregate's theta equals the weighted sum of the types' thetas "
      "(for one commodity the weighting cancels, so the two weightings differ only for aggregates)",
      ok_w, f"consistency {S['consistency_max_dev']:.1e}; 2023 weights {A['2023']['weights']}")

# ---- 4. depreciation against BEA's published rates
pub, imp = R["published_rates_pdf"], R["delta_software_implied"]
dv = max(abs(imp[k][y] / (pub[k] / (1 - pub[k] / 2)) - 1.0) for k in ("Prepackaged", "Custom", "Own-account") for y in YEARS)
check("C10 the implied depreciation rates (FA 2.4 over the average of FA 2.1) match BEA's published geometric rates for the "
      "single-type lines (software: prepackaged 0.55, custom and own-account 0.33, parsed from Tablecandtext.pdf) once the "
      "half-year convention is allowed for — steady state D/K = delta/(1 - delta/2) — within 3% every year",
      dv <= 0.03 and pub.get("Prepackaged") == 0.55 and pub.get("Custom") == 0.33, f"max dev {dv:.4f}; published {pub}")

# ---- 5. hours and per income
E = R["economy"]
hy = [y for y in YEARS if "total_hours_io_fte" in E[y]]
hc = max(abs(E[y]["total_hours_io_fte"] / E[y]["nipa_fte_hours_domestic"] - 1.0) for y in hy)
check("C11 hours: I-O compensation over NIPA compensation per FTE-hour, summed over industries, reproduces NIPA's domestic "
      "FTE x 2,080 within 2% every year 1998-2023 (the redefinitions move compensation between industries, not in total)",
      hc <= 0.02 and len(hy) == 26, f"max dev {hc:.4f}")
wv = R["wrel_validation"]
gap = {s: [wv[s][y]["fte_basis"] / wv[s][y]["thread_hours_basis"] - 1.0 for y in wv[s]] for s in wv}
check("C12 finding: the FTE x 2,080 basis puts the chain's relative wage above the lambda thread's hours-worked basis for "
      "narrow, medium and broad in every year, by 4-20% [threshold set after seeing the data: observed 4.8-19.1%] (a uniform "
      "full-time year ignores that schedules differ across industries); the per-income central values are therefore "
      "bridged to the hours-worked basis via the medium set",
      all(0.04 <= min(v) and max(v) <= 0.20 for v in gap.values()),
      {s: (round(min(v), 3), round(max(v), 3)) for s, v in gap.items()})
mono, n_rows = True, 0
for yr, P in R["per_income"].items():
    o = P["oews"]
    mono &= all(o[a] < o[b] for a, b in zip(PCTS, PCTS[1:]))
    for lab, rw in P["rows"].items():
        for basis, v in rw.items():
            if isinstance(v, dict):
                for kind, vals in v.items():
                    n_rows += 1
                    mono &= all(vals[a] < vals[b] for a, b in zip(PCTS, PCTS[1:]))
check("C13 per income: the OEWS all-occupations percentiles rise P10 -> P90 and so does every row of machine-chain hours per "
      "hour replaced (every year, chain, theta basis and wage basis)", mono and n_rows > 1000, f"{n_rows} rows")
fm = R["fred_mirror_check"]
fd = max(abs(v["fred_A033RC1A027NBEA"] / v["nipa_A033RC"] - 1.0) for v in fm.values())
check("C14 the NIPA flat file (the publisher's own) and FRED's mirror agree on compensation of employees (A033RC) within 0.1%",
      fd <= 1e-3 and len(fm) >= 3, f"max dev {fd:.1e}")
da, sa = D[f"{SET}_investment_weighted"], A["2017"]
dmax = max(abs(da[f"theta_{t}"] - sa[f"theta_{t}"]) for t in ("tot", "dom", "domp"))
check("C15 aggregation: at the 2017 detail benchmark the investment-weighted machinery + software set's theta (77 detail "
      "commodities) is within 0.02 of the summary value on tot, dom and domp [threshold set after seeing the data: observed "
      "0.011]", dmax <= 0.02,
      f"detail {({k: round(v, 4) for k, v in da.items() if k.startswith('theta')})}; summary "
      f"{({t: round(sa['theta_' + t], 4) for t in ('tot', 'dom', 'domp')})}")
ai = D["ai_compute_composite"]
ai_ok = abs(sum(ai["weights"].values()) - 1.0) < 1e-12 and \
    all(abs(ai[k] - sum(w * D["types"][c][k] for c, w in ai["weights"].items())) < 1e-12
        for k in ("theta_tot", "cfc_tot", "nos_tot", "cfc_tot_economywide", "nos_tot_economywide")) and \
    ai["component_range_theta_tot"][0] <= ai["theta_tot"] <= ai["component_range_theta_tot"][1] and \
    abs(ai["cfc_tot_economywide"] + ai["nos_tot_economywide"] - ai["gos_tot"]) < 1e-12 and \
    all(D["types"][c]["summary_parent_ratio_capped"] for c in ai["weights"] if c.startswith("334")) and \
    "artefact of the cap" in ai["cfc_nos_split_note"]
check("C16 the AI-compute composite is the stated weighted sum of its four detail components (weights sum to 1) and lies "
      "within their range; its 334 components inherit a capped ratio, so its central CFC/NOS split is labelled an artefact "
      "of the cap and the economy-wide-ratio variant (CFC + NOS = GOS) is reported beside", ai_ok,
      f"theta_tot {ai['theta_tot']:.4f}; CFC/NOS central {ai['cfc_tot']:.3f}/{ai['nos_tot']:.3f}, economy-wide ratio "
      f"{ai['cfc_tot_economywide']:.3f}/{ai['nos_tot_economywide']:.3f}")

# ---- 6. every headline number read back from the stored series
T, H = R["targets_for_macro"], R["headline"]
nr, eco, dag = R["nipa_ratios"], R["economy"], R["delta_aggregates"]
BC = R["bundles"][BUNDLE_C]
TA = T[f"a_theta_V_{SET}_investment_weighted"]
rng_src = {"narrow set (333, 334, 335), investment-weighted": AGG["narrow|investment"]["2023"]["theta_tot"],
           "equipment + IPP, all commodities": R["bundles"]["equipment_ipp"]["2023"]["theta_tot"],
           "equipment + IPP + private nonresidential structures, all commodities (calibration (c))": BC["2023"]["theta_tot"],
           "machinery + software set (investment-weighted)": A["2023"]["theta_tot"]}
pairs = [
    (TA["tot_2023"], A["2023"]["theta_tot"]), (TA["dom_2023"], A["2023"]["theta_dom"]),
    (TA["domp_2023"], A["2023"]["theta_domp"]),
    (TA["tot_mean_2015_2023"], M([A[y]["theta_tot"] for y in W15])),
    (TA["dom_mean_2015_2023"], M([A[y]["theta_dom"] for y in W15])),
    (TA["domp_mean_2015_2023"], M([A[y]["theta_domp"] for y in W15])),
    (TA["set_share_of_equipment_ipp_investment_2023"],
     eco["2023"][f"{SET}_investment_$m"] / eco["2023"]["inv_equipment_ipp_$m"]),
    (TA["range_2023_tot"][0], min(rng_src.values())), (TA["range_2023_tot"][1], max(rng_src.values())),
    (T["c_nipa"]["cfc_over_gdp"]["2023"], nr["2023"]["cfc_over_gdp"]),
    (T["c_nipa"]["nos_over_gdp"]["mean_2015_2023"], M([nr[y]["nos_over_gdp"] for y in W15])),
    (T["e_delta_equipment_ipp_private"]["2023"], dag["equipment + IPP (private nonresidential)"]["2023"]),
    (T["e_delta_equipment_ipp_private"][f"{SET}_investment_weighted_2023"], S["delta"]["2023"]),
    (T[f"traced_{SET}_2023"], S["traced"]["2023"]),
    (T[f"traced_service_central_{SET}_2023"], S["traced_service_central"]["2023"]),
    (H["theta_V_2023_tot"], A["2023"]["theta_tot"]), (H["theta_service_2023_central"], S["service"][RC]["2023"]),
    (H["ai_compute_theta_tot_2017"], ai["theta_tot"]), (H["gate_max_abs_dev"], g["max_abs_dev"]),
    (H["hours_bridge_2023"], R["hours_bridge_medium"]["2023"]),
]
pairs += [(TA["range_2023_tot_across_scopes"][k], v) for k, v in rng_src.items()]
for k, v in T["b_labour_embodied_in_fixed_investment"].items():
    pairs += [(v["2023"], eco["2023"][k]), (v["mean_2015_2023"], M([eco[y][k] for y in W15]))]
for k in ("theta", "cfc", "tax", "nos"):
    v = T[f"d_cost_split_{SET}_investment_weighted_tot"][k]
    pairs += [(v["2023"], A["2023"][f"{k}_tot"]), (v["mean_2015_2023"], M([A[y][f"{k}_tot"] for y in W15]))]
for blk in ("d_cost_split_dom", f"d_cost_split_{SET}_domp", f"d_bounds_{SET}"):
    for k, v in T[blk].items():
        pairs += [(v["2023"], A["2023"][k]), (v["mean_2015_2023"], M([A[y][k] for y in W15]))]
for r, v in T[f"service_flow_{SET}_2023"].items():
    pairs.append((v, S["service"][r]["2023"]))
for c, v in H["theta_2023_by_type_tot"].items():
    pairs.append((v, R["types"][c]["annual"]["2023"]["theta_tot"]))
pi = R["per_income"]["2023"]["rows"][f"{SET}|investment|tot"]
for key, basis, kind in (("per_income_2023_service_central", SVC, "central_relative_hours_worked"),
                         ("per_income_2023_service_relative_fte", SVC, "relative_fte"),
                         ("per_income_2023_asset", "asset", "central_relative_hours_worked"),
                         ("per_income_2023_service_traced", SVT, "central_relative_hours_worked")):
    pairs += [(H[key][p], pi[basis][kind][p]) for p in PCTS]
for k, v in T[f"g1_{BUNDLE_C}_full_chain"].items():
    pairs += [(v["2023"], BC["2023"][k]), (v["mean_2015_2023"], M([BC[y][k] for y in W15]))]
DU = R["import_convention_domestic_use"][BUNDLE_C]
for k, v in T[f"g1_{BUNDLE_C}_import_convention_domestic_use"].items():
    pairs += [(v["2023"], DU["2023"][k]), (v["mean_2015_2023"], M([DU[y][k] for y in W15]))]
for b in BASES:
    hr_ = sorted([BC["2023"][f"hours_share_of_total_hours_fte_{b}"], BC["2023"][f"hours_share_of_total_hours_bridged_{b}"]])
    pairs += [(T[f"g1_{BUNDLE_C}_hours_share_range_2023"][b][0], hr_[0]), (T[f"g1_{BUNDLE_C}_hours_share_range_2023"][b][1], hr_[1])]
for b in BASES:
    for q in ("theta", "cfc", "tax", "nos", "imports"):
        pairs += [(H["g1_bundle_2023"][b][q], BC["2023"][f"{q}_{b}"]),
                  (H["g1_bundle_mean_2015_2023"][b][q], M([BC[y][f"{q}_{b}"] for y in W15]))]
MV = R["model_vs_data"]
pairs += [(H["model_vs_data_multiples"][k], v["multiple"]) for k, v in MV["comparisons"].items()]
worst = max(abs(a - b) for a, b in pairs)
check("C17 every headline number and every target for the macro block (with the G1 bundle's split, its import-convention "
      "sensitivity and hours range, the set's range and the model-against-data multiples) is read back from the stored series "
      "chains.py wrote (means recomputed from the annual values)", worst <= 1e-12, f"{len(pairs)} numbers, max dev {worst:.1e}")
cs = max(abs(A[y]["theta_tot"] + A[y][f"cfc_tot{v}"] + A[y]["tax_tot"] + A[y][f"nos_tot{v}"] - 1.0)
         for y in YEARS for v in ("", "_uncapped", "_economywide"))
check("C18 the aggregate's cost split (labour / CFC / taxes / NOS) sums to 1 in every year and CFC variant", cs <= 1e-5,
      f"max dev {cs:.1e}")

# ---- 7. G1: the full-chain cost split of calibration (c)'s bundle, on every basis
bad = []
for bname, BB in R["bundles"].items():
    kinv = f"inv_{bname}"
    for y in YEARS:
        s = BB[y]
        for b in BASES:
            for v in ("", "_uncapped", "_economywide"):
                tot = s[f"theta_{b}"] + s[f"cfc_{b}{v}"] + s[f"tax_{b}"] + s[f"nos_{b}{v}"] + s[f"imports_{b}"]
                if abs(tot - 1.0) > 1e-5:
                    bad.append((bname, y, b, v, "adds to", tot))
            if abs(s[f"theta_{b}"] - eco[y][f"{kinv}_theta_{b}"]) > 1e-12:
                bad.append((bname, y, b, "theta vs economy"))
            if s[f"comp_share_of_total_comp_{b}"] != eco[y][f"{kinv}_comp_{b}_share_of_total_comp"]:
                bad.append((bname, y, b, "comp share"))
            if y != "1997" and (s[f"hours_share_of_total_hours_fte_{b}"] != eco[y][f"{kinv}_hours_{b}_share_of_total_hours"]
                                or s[f"hours_share_of_total_hours_bridged_{b}"]
                                != eco[y][f"{kinv}_hours_{b}_share_of_total_hours_bridged_hours_worked"]):
                bad.append((bname, y, b, "hours share"))
            if not (s[f"sigma_upper_closed_{b}"] == s[f"nos_{b}"]
                    and s[f"sigma_upper_open_{b}"] == s[f"nos_{b}"] + s[f"imports_{b}"]
                    and abs(s[f"multiplier_lower_bound_closed_{b}"] - (1 + s[f"theta_{b}"] / s[f"nos_{b}"])) < 1e-12
                    and abs(s[f"multiplier_lower_bound_open_{b}"]
                            - (1 + s[f"theta_{b}"] / (s[f"nos_{b}"] + s[f"imports_{b}"]))) < 1e-12
                    and abs(s[f"multiplier_floor_any_sigma_{b}"] - 1 / (1 - s[f"theta_{b}"])) < 1e-12
                    and s[f"multiplier_floor_any_sigma_{b}"] <= s[f"multiplier_lower_bound_open_{b}"] + 1e-12
                    <= s[f"multiplier_lower_bound_closed_{b}"] + 2e-12):
                bad.append((bname, y, b, "bounds"))
        if not (s["imports_tot"] == 0.0 and abs(s["imports_domp"] - s["imports_intermediate_domp"] - s["imports_final_domp"]) < 1e-15
                and s["theta_domp"] <= s["theta_dom"] + 1e-4 <= s["theta_tot"] + 2e-4):
            bad.append((bname, y, "imports / ordering"))
for y in YEARS:                      # the set's bounds (targets d_bounds_*) by the same formulas
    s = A[y]
    for b in BASES:
        if not (s[f"sigma_upper_open_{b}"] == s[f"nos_{b}"] + s[f"imports_{b}"]
                and abs(s[f"multiplier_lower_bound_closed_{b}"] - (1 + s[f"theta_{b}"] / s[f"nos_{b}"])) < 1e-12
                and abs(s[f"multiplier_lower_bound_open_{b}"]
                        - (1 + s[f"theta_{b}"] / (s[f"nos_{b}"] + s[f"imports_{b}"]))) < 1e-12
                and abs(s[f"theta_{b}"] + s[f"cfc_{b}"] + s[f"tax_{b}"] + s[f"nos_{b}"] + s[f"imports_{b}"] - 1) < 1e-5):
            bad.append((SET, y, b, "bounds"))
for y in YEARS:                      # the import-convention sensitivity (domestic-use shares) for the (c) bundle
    s, s0 = DU[y], BC[y]
    for b in ("dom", "domp"):
        if abs(s[f"theta_{b}"] + s[f"cfc_{b}"] + s[f"tax_{b}"] + s[f"nos_{b}"] + s[f"imports_{b}"] - 1) > 1e-5 or \
                s[f"sigma_upper_open_{b}"] != s[f"nos_{b}"] + s[f"imports_{b}"] or \
                abs(s[f"multiplier_lower_bound_open_{b}"] - (1 + s[f"theta_{b}"] / s[f"sigma_upper_open_{b}"])) > 1e-12 or \
                abs(s["imports_domp"] - s["imports_intermediate_domp"] - s["imports_final_domp"]) > 1e-15:
            bad.append(("domestic-use", y, b))
    if not (s["theta_domp"] < s0["theta_domp"] and s["imports_domp"] > s0["imports_domp"]):
        bad.append(("domestic-use", y, "direction"))
scope_txt = T[f"g1_{BUNDLE_C}_scope"]
named = "q/(q + M)" in scope_txt and "1 - M/(q + M - X)" in scope_txt and "borrows the medium" in scope_txt \
    and "q/(q + M)" in R["meta"]["import_convention"]
du_gap = [BC[y]["theta_domp"] - DU[y]["theta_domp"] for y in YEARS]
# the bundle is exactly what calibration (c) reads: its theta and chain shares are the economy keys calibration_c.py takes
with open(os.path.join(ROOT, "results", "calibration_c.json"), encoding="utf-8") as fh:
    cread = json.load(fh)["targets"]["from_chains_json"]
same_c = abs(cread["theta_V_tot"]["value"] - BC["2023"]["theta_tot"]) < 1e-12 and \
    abs(cread["chain_share_tot"]["value"] - BC["2023"]["comp_share_of_total_comp_tot"]) < 1e-12 and \
    abs(cread["theta_V_domp"]["value"] - BC["2023"]["theta_domp"]) < 1e-12 and \
    abs(cread["chain_share_domp"]["value"] - BC["2023"]["comp_share_of_total_comp_domp"]) < 1e-12
b23 = {b: {q: round(BC["2023"][f"{q}_{b}"], 4) for q in ("theta", "cfc", "tax", "nos", "imports")} for b in BASES}
check("C19 G1: calibration (c)'s bundle (equipment + IPP + private nonresidential structures, all commodities) has its full "
      "chain split on tot, dom and domp: compensation + CFC + taxes + NOS + imported content = 1 in every year, basis, "
      "bundle and CFC variant (tolerance 1e-5, the published TR); its theta, compensation and hours shares are the "
      "economy keys calibration_c.py reads (and the values recorded in results/calibration_c.json); sigma_upper = NOS "
      "(closed) and NOS + imports (open); 1/(1 - theta) <= open bound <= closed bound. The dom/domp import convention (q/(q + M), "
      "the thread's) is named, with the domestic-use sensitivity (1 - M/(q + M - X)) stored for the bundle: its split sums to 1 "
      "on dom and domp every year and it lowers theta_domp and raises imported content every year; the scope note also says "
      "the bundle's bridged hours borrow the medium set's bridge",
      not bad and same_c and named, f"violations {bad[:4]}; 2023 {b23}; domestic-use convention lowers theta_domp by "
      f"{min(du_gap):.4f}-{max(du_gap):.4f} (2023 {BC['2023']['theta_domp']:.4f} -> {DU['2023']['theta_domp']:.4f}; open multiplier "
      f"bound {BC['2023']['multiplier_lower_bound_open_domp']:.2f} -> {DU['2023']['multiplier_lower_bound_open_domp']:.2f})")


# ---- 8. G2: an independent path from the raw files (own readers; neither the lambda engine nor chains.py)
def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0                                     # '...' (not shown: zero) and blanks


def io_sheet(member_zip, member, yr):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(zipfile.ZipFile(os.path.join(CACHE, member_zip)).read(member)),
                                read_only=True, data_only=True)
    rows_ = list(wb[str(yr)].iter_rows(values_only=True))
    h = next(i for i, r in enumerate(rows_) if r[0] in ("Code", "IOCode"))
    codes = [None if c is None else str(c).strip() for c in rows_[h - 1]]
    names = [None if c is None else str(c).strip() for c in rows_[h]]
    body = {}
    for r in rows_[h + 1:]:
        if r[0] is not None:
            body[str(r[0]).strip()] = {codes[j]: num(r[j]) for j in range(2, len(r)) if codes[j]}
    return body, {codes[j]: names[j] for j in range(2, len(codes)) if codes[j]}


def line_sheet(path, sheet):
    import openpyxl
    rows_ = list(openpyxl.load_workbook(path, read_only=True, data_only=True)[sheet].iter_rows(values_only=True))
    h = next(i for i, r in enumerate(rows_) if r[0] == "Line")
    yc = {int(str(c)[:4]): j for j, c in enumerate(rows_[h]) if c is not None and re.fullmatch(r"\d{4}(\.0)?", str(c))}
    return [(re.sub(r"\\\d+\\", "", str(r[1])).strip(), {y: num(r[j]) for y, j in yc.items()})
            for r in rows_[h + 1:] if r[0] is not None and str(r[0]).strip().isdigit()]


INV = ["F02E", "F02N", "F06E", "F06N", "F07E", "F07N", "F10E", "F10N"]
MEMBERS = ["333", "334", "335", "511", "514", "5415", "532RL"]
ind_path = {}
for yr in (1997, 2017, 2023):
    TRb, _ = io_sheet("AllTablesSUP.zip", "IxC_TR_1997-2023_Summary.xlsx", yr)
    DRb, io_names = io_sheet("AllTablesIO.zip", "CxI_DR_1997-2023_Summary.xlsx", yr)
    USb, _ = io_sheet("AllTablesIO.zip", "IOUse_After_Redefinitions_PRO_1997-2023_Summary.xlsx", yr)
    inds = [i for i in TRb if i in DRb["V001"]]
    coms = list(next(iter(TRb.values())).keys())
    tr = np.array([[TRb[i][c] for c in coms] for i in inds])
    vw, vg = np.array([DRb["V001"][i] for i in inds]), np.array([DRb["V003"][i] for i in inds])
    theta_c = vw @ tr                                    # full-chain compensation per $ of each commodity (tot)
    wset = np.array([max(sum(USb.get(c, {}).get(k, 0.0) for k in INV), 0.0) if c in MEMBERS else 0.0 for c in coms])
    wbun = np.array([max(sum(USb.get(c, {}).get(k, 0.0) for k in INV + ["F02S"]), 0.0) for c in coms])
    wset, wbun = wset / wset.sum(), wbun / wbun.sum()
    ind_path[yr] = {"set_theta_tot": float(theta_c @ wset), "bundle_theta_tot": float(theta_c @ wbun)}
    if yr == 2023:                                        # the central CFC share, with its own industry map
        fa = {}
        for n, v in line_sheet(os.path.join(CACHE, "Section3All_xls.xlsx"), "FAAt304ESI-A"):
            fa.setdefault(n, v)
        fa24 = dict(line_sheet(os.path.join(CACHE, "Section2All_xls.xlsx"), "FAAt204-A")[::-1])
        va_rows = line_sheet(os.path.join(CACHE, "ValueAdded.xlsx"), "TVA113-A")
        gos, cur, seen = {}, None, {}
        for n, v in va_rows:
            if n == "Gross operating surplus" and cur is not None:
                gos.setdefault(cur, []).append(v[yr])
            elif n not in ("Compensation of employees", "Taxes on production and imports less subsidies"):
                cur = n
        flat = {}
        with open(os.path.join(CACHE, "NipaDataA.txt"), encoding="utf-8") as fh:
            for r in csv.reader(fh):
                if r and r[0] in ("B833RC", "B846RC") and r[1] == str(yr):
                    flat[r[0]] = float(r[2].replace(",", ""))
        FA_NAME = {"Rail transportation": ["Railroad transportation"],
                   "Publishing industries, except internet (includes software)": ["Publishing industries (includes software)"],
                   "Data processing, internet publishing, and other information services":
                       ["Information and data processing services"],
                   "Federal Reserve banks, credit intermediation, and related activities":
                       ["Federal Reserve banks", "Credit intermediation and related activities"]}
        resid = fa24["Residential equipment"][yr] + fa24["Residential structures"][yr]
        ratio, fallback = {}, []
        for i in inds:
            nm = io_names[i]
            if "general government" in nm:
                ratio[i] = 1.0
                continue
            if nm in ("Federal government enterprises", "State and local government enterprises"):
                k = 0 if nm.startswith("Federal") else 1
                c_, g_ = flat["B833RC" if k == 0 else "B846RC"], gos["Government enterprises"][k]
            elif nm == "Housing":
                c_, g_ = resid, gos["Housing"][0]
            elif nm == "Other real estate":
                c_, g_ = fa["Real estate"][yr] - resid, gos["Other real estate"][0]
            else:
                c_, g_ = sum(fa[n][yr] for n in FA_NAME.get(nm, [nm])), gos[nm][0]
            if g_ <= 0:
                fallback.append(i)
            ratio[i] = c_ / g_
        vcfc = np.array([min(ratio[i], 1.0) * vg[j] for j, i in enumerate(inds)])
        ind_path[yr]["bundle_cfc_tot"] = float(vcfc @ tr @ wbun)
        ind_path[yr]["set_cfc_tot"] = float(vcfc @ tr @ wset)
        ind_path[yr]["fallback"] = fallback
        # route S (domp, the basis calibration (c)'s central reads): the Leontief system self-built from the Use and Make
        # levels (no published TR or DR), the domestic purge from Use's imports and exports, both import conventions
        MKb, _ = io_sheet("AllTablesIO.zip", "IOMake_After_Redefinitions_PRO_1997-2023_Summary.xlsx", yr)
        Vmk = np.array([[MKb[i].get(c, 0.0) for c in coms] for i in inds])
        g_s, q_s = Vmk.sum(axis=1), Vmk.sum(axis=0)
        B_s = np.array([[USb[c][i] / g_s[k] for k, i in enumerate(inds)] for c in coms])
        W_s = Vmk / np.where(q_s > 0, q_s, 1.0)[None, :]
        tr_s = W_s @ np.linalg.inv(np.eye(len(coms)) - B_s @ W_s)
        vw_s = np.array([USb["V001"][i] for i in inds]) / g_s
        vg_s = np.array([USb["V003"][i] for i in inds]) / g_s
        vcfc_s = np.array([min(ratio[i], 1.0) * vg_s[j] for j, i in enumerate(inds)])
        imp_s = np.array([max(-USb[c].get("F050", 0.0), 0.0) for c in coms])
        exp_s = np.array([USb[c].get("F040", 0.0) for c in coms])
        fl_s = np.array([max(sum(USb.get(c, {}).get(k, 0.0) for k in INV + ["F02S"]), 0.0) for c in coms])
        f_s = fl_s / fl_s.sum()
        ph_t = np.clip(q_s / np.where(q_s + imp_s > 0, q_s + imp_s, 1.0), 0.0, 1.0)
        den_s = q_s + imp_s - exp_s
        ph_d = np.where(den_s > 0, 1.0 - np.clip(imp_s / np.where(den_s > 0, den_s, 1.0), 0.0, 1.0), ph_t)
        comp_total_s = sum(USb["V001"][i] for i in inds)
        ind_path[yr]["tr_self_built_vs_published"] = float(np.max(np.abs(tr_s - tr)))
        for conv, ph in (("thread", ph_t), ("domestic_use", ph_d)):
            trd = W_s @ np.linalg.inv(np.eye(len(coms)) - (ph[:, None] * B_s) @ W_s)
            rp = trd @ (ph * f_s)
            ind_path[yr][f"{conv}_theta_domp"] = float(vw_s @ rp)
            ind_path[yr][f"{conv}_nos_domp"] = float(vg_s @ rp - vcfc_s @ rp)
            ind_path[yr][f"{conv}_imports_domp"] = float(((1 - ph)[:, None] * B_s).sum(axis=0) @ rp + ((1 - ph) * f_s).sum())
            ind_path[yr][f"{conv}_comp_share_of_total_comp_domp"] = float(vw_s @ trd @ (ph * fl_s)) / comp_total_s
stored = {yr: {"set_theta_tot": A[str(yr)]["theta_tot"], "bundle_theta_tot": BC[str(yr)]["theta_tot"]} for yr in ind_path}
stored[2023].update(bundle_cfc_tot=BC["2023"]["cfc_tot"], set_cfc_tot=A["2023"]["cfc_tot"])
dev2 = max(abs(ind_path[yr][k] - v) for yr in stored for k, v in stored[yr].items())
stored_s = {f"{conv}_{q}": src[q] for conv, src in (("thread", BC["2023"]), ("domestic_use", DU["2023"]))
            for q in ("theta_domp", "nos_domp", "imports_domp", "comp_share_of_total_comp_domp")}
dev_s = max(abs(ind_path[2023][k] - v) for k, v in stored_s.items())
check("C20 G2 independent path: from the raw BEA files with this check's own readers (openpyxl on the zipped I-O Use, Make, IxC "
      "TR and DR summaries; FA 3.4ESI and 2.4; GDP-by-industry TVA113; the NIPA flat file) and its own industry map (the "
      "I-O's industry names matched to GDP by industry, four FA spellings by hand), the investment-weighted theta_tot of "
      "the machinery + software set and of the (c) bundle (1997, 2017, 2023) and the central CFC share of both (2023) "
      "equal the stored values to <= 1e-9; and on the domp basis calibration (c)'s central reads, the (c) bundle's theta, "
      "NOS, imported content and chain share of compensation (2023), from a Leontief system self-built from the Use and Make "
      "levels (no published TR or DR), equal the stored values to <= 1e-5 under the thread's import convention and under the "
      "domestic-use sensitivity [1e-5: the published DR's rounding; see the header]",
      dev2 <= 1e-9 and not ind_path[2023]["fallback"] and dev_s <= 1e-5,
      f"max dev {dev2:.1e}; 2023 set theta {ind_path[2023]['set_theta_tot']:.5f}, bundle theta "
      f"{ind_path[2023]['bundle_theta_tot']:.5f}, bundle CFC {ind_path[2023]['bundle_cfc_tot']:.5f}, set CFC "
      f"{ind_path[2023]['set_cfc_tot']:.5f}; self-built TR vs published {ind_path[2023]['tr_self_built_vs_published']:.1e}; domp "
      f"max dev {dev_s:.1e}: theta {ind_path[2023]['thread_theta_domp']:.5f}, NOS {ind_path[2023]['thread_nos_domp']:.5f}, imports "
      f"{ind_path[2023]['thread_imports_domp']:.5f}, chain share {ind_path[2023]['thread_comp_share_of_total_comp_domp']:.5f} "
      f"(domestic-use: {ind_path[2023]['domestic_use_theta_domp']:.5f}, {ind_path[2023]['domestic_use_nos_domp']:.5f}, "
      f"{ind_path[2023]['domestic_use_imports_domp']:.5f}, {ind_path[2023]['domestic_use_comp_share_of_total_comp_domp']:.5f})")

# ---- 9. G7: the model against the data, from macro.py's own calibrations
mod_ok, mdev = True, 0.0
for cal, base in (("a", m.CALIBRATION_BASE), ("b", m.CAPITAL_BASE)):
    e, q, _, err = m.calibrate(base=base)
    u, d = e.delta / (e.rho + e.delta), e.a * e.delta
    want = {"theta_V_asset": q["theta_V"], "theta_V_service": q["theta_V"] * u, "theta_V_traced_asset": q["theta_V"] / (1 - d),
            "theta_V_traced_service": q["theta_V"] * u / (1 - d), "chain_share": q["chain_share"], "multiplier": q["multiplier"]}
    mdev = max(mdev, max(abs(MV["model"][cal][k] - v) / abs(v) for k, v in want.items()))
    mod_ok &= err <= 1e-8
data_src = {"set_asset": A["2023"]["theta_tot"], "set_service": S["service"][RC]["2023"],
            "set_traced_asset": S["traced"]["2023"], "set_traced_service": S["traced_service_central"]["2023"],
            "bundle_asset": BC["2023"]["theta_tot"],
            "bundle_comp_share": eco["2023"][f"inv_{BUNDLE_C}_comp_tot_share_of_total_comp"],
            "bundle_hours_share": eco["2023"][f"inv_{BUNDLE_C}_hours_tot_share_of_total_hours_bridged_hours_worked"],
            "bundle_hours_share_fte": eco["2023"][f"inv_{BUNDLE_C}_hours_tot_share_of_total_hours"],
            "eqipp_comp_share": eco["2023"]["inv_equipment_ipp_comp_tot_share_of_total_comp"],
            "set_comp_share": A["2023"]["theta_tot"] * eco["2023"][f"{SET}_investment_$m"] / eco["2023"]["total_comp_io"]}
SCOPE = {"set": "machinery + software set (investment-weighted)", "bundle": "calibration (c)'s bundle",
         "eqipp": "equipment + IPP (private and government)"}
for cal in ("a", "b"):
    mm = MV["model"][cal]
    data_src[f"set_service_at_model_{cal}_rho_delta"] = A["2023"]["theta_tot"] * mm["delta"] / (mm["rho"] + mm["delta"])
cdev = max(max(abs(v["data"] - data_src[v["data_key"]]), abs(v["multiple"] - v["data"] / MV["model"][v["calibration"]][v["model_key"]]))
           for v in MV["comparisons"].values())
C_ = MV["comparisons"]
ident = abs(C_["b|service_at_model_rho_delta"]["multiple"] - C_["b|asset"]["multiple"]) < 1e-9 and \
    abs(C_["a|service_at_model_rho_delta"]["multiple"] - C_["a|asset"]["multiple"]) < 1e-9 and \
    MV["model"]["a"]["rho"] == 0.0 and MV["model"]["a"]["delta"] == 1.0 and \
    R["meta"]["rho_central"] == m.CAPITAL_BASE.rho and \
    not C_["a|service_data_rho_central"]["like_for_like"] and \
    all(v["scope"].startswith(SCOPE[v["data_key"].split("_")[0]]) for v in C_.values()) and \
    {v["data_key"] for v in C_.values() if v["model_key"] == "chain_share"} == {"bundle_comp_share", "bundle_hours_share",
                                                                            "bundle_hours_share_fte", "eqipp_comp_share", "set_comp_share"}
check("C21 G7 the model against the data: calibrations (a) and (b) recomputed here from macro.py reproduce the stored model "
      "thetas (asset, service, traced) and chain shares; every data value is the stored series it names and every multiple "
      "is data/model; on the model's own (rho, delta) the service multiple equals the asset multiple; (a)'s data-at-rho-"
      "0.0575 figure is marked mixed-basis; rho_central is macro.CAPITAL_BASE.rho; every row names its data scope, and the "
      "chain share is set against three scopes (the set, equipment + IPP, (c)'s bundle) and both hours measures",
      mod_ok and mdev <= 1e-9 and cdev <= 1e-12 and ident,
      f"model rel dev {mdev:.1e}; multiples dev {cdev:.1e}; like-for-like: "
      + ", ".join(f"{k} {v['multiple']:.0f}x" for k, v in C_.items() if v["like_for_like"]))

# ---- 10. G3: the capped and fallback ledger, chain-weighted
bad3 = []
for key, agg_row in ((SET, lambda y: A[y]), (BUNDLE_C, lambda y: BC[y])):
    for y in YEARS:
        cw = led[y]["chain_weighted"][key]
        if abs(cw["cap_removed_total"] - (agg_row(y)["cfc_tot_uncapped"] - agg_row(y)["cfc_tot"])) > 1e-12:
            bad3.append((key, y, "removal sum"))
        if sorted(cw["capped"]) != sorted(led[y]["cap_binds"]) or sorted(cw["fallback"]) != sorted(led[y]["per_output_fallback"]):
            bad3.append((key, y, "listing"))
    for i, e_ in CF[key]["industries"].items():
        if e_["capped_years"] != [int(y) for y in YEARS if i in led[y]["cap_binds"]] or \
                e_["fallback_years"] != [int(y) for y in YEARS if i in led[y]["per_output_fallback"]]:
            bad3.append((key, i, "years"))
    if set(CF[key]["industries"]) != {i for y in YEARS for i in led[y]["cap_binds"] + led[y]["per_output_fallback"]}:
        bad3.append((key, "industry set"))
c5415 = CF[SET]["industries"]["5415"]
check("C22 G3 every capped and every fallback industry-year is listed with its chain-weighted part (for the set and the (c) "
      "bundle); the removals sum to uncapped less central CFC every year; 5415 is capped in 1999 and 2001-2007 and on the "
      "fallback in 2000",
      not bad3 and c5415["capped_years"] == [1999] + list(range(2001, 2008)) and c5415["fallback_years"] == [2000],
      f"violations {bad3[:4]}; cap removes at most {CF[SET]['cap_removed_total_max']:.4f} (set), "
      f"{CF[BUNDLE_C]['cap_removed_total_max']:.4f} (bundle) of $1; 2015-2023 {CF[SET]['cap_removed_total_max_2015_2023']:.4f}"
      f" / {CF[BUNDLE_C]['cap_removed_total_max_2015_2023']:.4f}")

# ---- 11. G5 variants, G6 co-headline, G8 aliases, vintage
rp = R["service_relative_price_variant"]
w23 = A["2023"]["weights"]
dl = {c: R["types"][c]["delta"]["2023"] for c in TYPES}
rc = R["meta"]["rho_central"]
svc_rp = sum(w * R["types"][k]["annual"]["2023"]["theta_tot"] for k, w in w23.items()) / \
    sum(w * (rc + dl[k] - rp["pi_rel_mean"][k]) / dl[k] for k, w in w23.items() if w > 0)
pv = R["published_rate_variant_2023"]["Prepackaged"]
g5 = abs(svc_rp - rp["set_2023"]["service_with_relative_price"]) < 1e-12 and \
    abs(rp["set_2023"]["service_central"] - S["service"][RC]["2023"]) < 1e-12 and \
    all(rp["types_2023"][c]["service_with_relative_price"] < rp["types_2023"][c]["service_central"]
        for c in TYPES if rp["pi_rel_mean"][c] < 0) and \
    abs(pv["service_over_asset_published"] - 0.55 / (rc + 0.55)) < 1e-12 and \
    abs(pv["service_over_asset_implied"] - pv["implied_D_over_Kbar_2023"] / (rc + pv["implied_D_over_Kbar_2023"])) < 1e-12
bridge_id = max(abs(pi[SVC]["relative_fte"][p] / pi[SVC]["central_relative_hours_worked"][p] - R["hours_bridge_medium"]["2023"])
                for p in PCTS)
def keys_with(o, needle, path=""):
    """Every key path in o whose key contains needle."""
    out = []
    if isinstance(o, dict):
        for k, v in o.items():
            if needle in str(k):
                out.append(path + "/" + str(k))
            out += keys_with(v, needle, path + "/" + str(k))
    elif isinstance(o, list):
        for i, v in enumerate(o):
            out += keys_with(v, needle, f"{path}[{i}]")
    return out


RN = R["meta"]["renamed_keys"]
old_left = [k for k in keys_with(R, "all_machinery") if not k.startswith("/meta/renamed_keys")]
alias_ok = not old_left and "aliases" not in R["meta"] and f"{SET}|investment" in AGG and \
    f"a_theta_V_{SET}_investment_weighted" in T and f"d_cost_split_{SET}_investment_weighted_tot" in T and \
    all(f"{SET}|investment|{b}" in R["per_income"][y]["rows"] for y in R["per_income"] for b in BASES) and \
    len([k for k in RN if k != "note"]) == 4
VT = R["meta"]["vintage"]
vint_ok = all("sha256_16" in v for k, v in VT.items() if k != "OEWS national zips") and \
    len(VT["OEWS national zips"]) == 25 and \
    all(any(s_.startswith("File created") for s_ in VT[f]["stamps"][sh]) for f in VT if "stamps" in VT[f] for sh in VT[f]["stamps"]) and \
    all(any(s_.startswith("Data published") for s_ in VT[f]["stamps"][sh])
        for f in ("ValueAdded.xlsx", "GrossOutput.xlsx", "Section6All_xls.xlsx") for sh in VT[f]["stamps"])
tc = R["meta"]["oews_top_code_hourly"]
check("C23 G5, G6, G8 and the vintage: the relative-price user-cost variant of the set's service share is recomputed from "
      "the stored pieces and lies below the central one wherever the price falls; the published-rate variant is the stated "
      "formula; the relative_fte co-headline is the central value over the hours bridge; the old 'all_machinery' copies are "
      "gone (no reader; meta.renamed_keys maps the four old names to keys that exist); every source file carries its size, hash and own "
      "stamp lines; the OEWS top code is read from each year's field descriptions ($100 in May 2021, $115 from May 2022)",
      g5 and bridge_id < 1e-12 and alias_ok and vint_ok and tc.get("2021") == 100.0 and tc.get("2022") == 115.0,
      f"set service 2023 {rp['set_2023']['service_central']:.4f} -> {rp['set_2023']['service_with_relative_price']:.4f} "
      f"with relative prices; bridge identity {bridge_id:.1e}; FA {VT['Section3All_xls.xlsx']['stamps']['FAAt304ESI-A']}, "
      f"GDPxInd {VT['ValueAdded.xlsx']['stamps']['TVA113-A']}")

n_ok = sum(ok for _, ok, _ in RESULTS)
if n_ok == len(RESULTS):
    print(f"\nALL GREEN ({n_ok})")
else:
    print(f"\n{len(RESULTS) - n_ok} FAILED of {len(RESULTS)}: " + ", ".join(n.split()[0] for n, ok, _ in RESULTS if not ok))
sys.exit(0 if n_ok == len(RESULTS) else 1)
